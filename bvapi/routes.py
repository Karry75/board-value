# -*- coding: utf-8 -*-
"""网点价值与财务收支看板 - API 路由"""
import sys, os, io, datetime, time
sys.path.insert(0, r"D:\Marvis K\janus\board-value")
sys.path.insert(0, r"D:\Marvis K\janus\electric")

from flask import Blueprint, request, jsonify, send_file
from bvcore import metrics

bp = Blueprint('api', __name__)

def _int(v, d=0):
    try:
        return int(v)
    except Exception:
        return d

def _str(v, d=''):
    return v if v is not None else d

def ok(data, **extra):
    r = {'code': 0, 'data': data}
    r.update(extra)
    return jsonify(r)

def err(msg):
    return jsonify({'code': 1, 'msg': str(msg)})

def _export_excel(headers, rows, sheet_name='明细', filename='export.xlsx'):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    # 表头
    fill = PatternFill('solid', fgColor='1F4E79')
    font = Font(color='FFFFFF', bold=True)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for ri, row in enumerate(rows, 2):
        for ci, v in enumerate(row, 1):
            ws.cell(ri, ci, v)
    # 列宽
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(30, (len(str(headers[c - 1])) + 4)))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ---------------- 筛选维度 ----------------
@bp.route('/filters')
def filters():
    if not metrics._db_ready():
        hit = metrics._cache_load('filters')
        if hit:
            hit['data']['data_source'] = 'cache'
            return ok(hit['data'])
        return ok(metrics.demo.filters())
    amap = metrics.load_agency_map()
    agencies = [{'id': k, 'name': v.get('name') or ''} for k, v in sorted(amap.items()) if v.get('is_del', 0) == 0 or 'is_del' not in v]
    cities = db_cities()
    products = [{'id': k, 'name': v} for k, v in sorted(metrics.load_battery_products().items())]
    data = {'agencies': agencies, 'cities': cities, 'products': products,
            'data_source': 'live',
            'data_date': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
    metrics._cache_save('filters', {'data': data, 'saved_at': time.time()})
    return ok(data)

def db_cities():
    try:
        rows = metrics.db.query("SELECT DISTINCT city FROM t_site WHERE is_del=0 AND city IS NOT NULL AND city<>'' ORDER BY city")
        return [r['city'] for r in rows]
    except Exception:
        return []

# ---------------- 网点价值明细 ----------------
@bp.route('/site_value')
def site_value():
    try:
        agency_id = _int(request.args.get('agency_id'))
        city = _str(request.args.get('city'))
        area = _str(request.args.get('area'))
        keyword = _str(request.args.get('keyword'))
        battery_product = _int(request.args.get('battery_product'))
        site_status = _str(request.args.get('site_status'))
        merchant_key = _str(request.args.get('merchant_key'))
        agency_key = _str(request.args.get('agency_key'))
        date_from = _str(request.args.get('date_from')) or None
        date_to = _str(request.args.get('date_to')) or None
        page = max(1, _int(request.args.get('page'), 1))
        page_size = min(200, max(10, _int(request.args.get('page_size'), 50)))
        rows, df, dt = metrics.build_site_value(
            agency_id=agency_id or None, city=city or None, area=area or None,
            keyword=keyword or None, battery_product=battery_product or None,
            site_status=site_status or None, date_from=date_from, date_to=date_to,
            merchant_key=merchant_key or None, agency_key=agency_key or None)
        total = len(rows)
        page_rows = rows[(page - 1) * page_size: page * page_size]
        return ok(page_rows, total=total, page=page, page_size=page_size,
                  date_from=df, date_to=dt, demo=metrics._LAST_SOURCE == 'demo',
                  data_source=metrics._LAST_SOURCE,
                  data_date=datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
    except Exception as e:
        return err(e)

@bp.route('/site_value/export')
def site_value_export():
    try:
        agency_id = _int(request.args.get('agency_id'))
        city = _str(request.args.get('city'))
        area = _str(request.args.get('area'))
        keyword = _str(request.args.get('keyword'))
        battery_product = _int(request.args.get('battery_product'))
        site_status = _str(request.args.get('site_status'))
        merchant_key = _str(request.args.get('merchant_key'))
        agency_key = _str(request.args.get('agency_key'))
        date_from = _str(request.args.get('date_from')) or None
        date_to = _str(request.args.get('date_to')) or None
        rows, df, dt = metrics.build_site_value(
            agency_id=agency_id or None, city=city or None, area=area or None,
            keyword=keyword or None, battery_product=battery_product or None,
            site_status=site_status or None, date_from=date_from, date_to=date_to,
            merchant_key=merchant_key or None, agency_key=agency_key or None)
        headers = ['数据日期起', '数据日期止', '城市', '区域', '街道', '代理商ID', '代理商名称', '代理商等级',
                   '网点ID', '网点名称', '网点状态', '商户ID', '商户名称', '创客手机号', '法人手机号', '店长',
                   '电池产品', '电表类型', '电费结算方式', '结算周期', '电费单价(元/度)', '电费金额(元)', '电费笔数',
                   '系统结算电费(元)', '线下对公电费(元)', '供电局划扣(元)',
                   '电费明细', '分成笔数', '分成金额(元)', '公司支付分成(元)', '网点退回(元)', '分成比例', '分成明细',
                   '网点用户数', '换电次数', '换电金额(元)', '换电量(Wh)', '销售笔数', '销售额(元)', '退款金额(元)',
                   '赠送包数', '赠送包估值(元)', '激励补贴(元)', '业务员提成(元)',
                   '设备柜数', '设备电池数', '设备投入(元)', '设备月分摊(元)',
                   '产出合计(元)', '投入合计(元)', '毛利(元)', '毛利率(%)', '净利润(元)', '产出投入比', '综合评分', '价值等级']
        data = []
        for r in rows:
            share_detail_txt = '; '.join('%s(%s%s,%.2f元)' % (it['expense_name'], it['expense_value_type'], ('%s' % it['expense_value'] if it['expense_value'] is not None else ''), it['fee_y']) for it in r.get('share_detail', []))
            share_ratio_txt = '; '.join('%s:%s%s' % (it['expense_name'], it['expense_value'], it['expense_value_type']) for it in r.get('share_detail', []) if it.get('expense_value_type') and '比例' in it['expense_value_type'])
            data.append([r['date_from'], r['date_to'], r['city'], r['area'], r['street'], r['agency_id'], r['agency_name'], r['agency_level'],
                         r['site_id'], r['site_name'], r['site_status'], r['merchant_id'], r['merchant_name'], r['maker_phone'], r['lp_phone'], r['store_manager_name'],
                         r['battery_product'], r['meter_type'], r['settle_way'], r['settle_cycle'], r['electric_price_y'], r['electric_fee_y'], r['electric_cnt'],
                         r.get('sys_elec_y', 0), r.get('offline_transfer_y', 0), r.get('gdj_deduct_y', 0),
                         '; '.join('%s:%s笔%.2f元' % (k, v['cnt'], v['fee_y']) for k, v in r['electric_detail'].items()),
                         r['share_cnt'], r['share_fee_y'], r['share_fee_in_y'], r['share_fee_out_y'], share_ratio_txt, share_detail_txt,
                         r.get('user_cnt', 0),
                         r['order_cnt'], r['order_fee_y'], r['use_power_wh'], r['sales_cnt'], r['sales_fee_y'], r['refund_fee_y'],
                         r['present_cnt'], r['present_fee_y'], r.get('incentive_y', 0), r.get('promoter_fee_y', 0),
                         r.get('device_cabinet_cnt', 0), r.get('device_battery_cnt', 0), r.get('device_cost_total_y', 0), r.get('device_cost_monthly_y', 0),
                         r['income_y'], r['cost_y'], r['profit_y'], r.get('profit_margin', 0), r.get('net_profit_y', r['profit_y']), r['ratio'], r['score'], r['level']])
        return _export_excel(headers, data, '网点价值明细', '网点价值明细_%s_%s.xlsx' % (df, dt))
    except Exception as e:
        return err(e)

# ---------------- 财务收支汇总(可视化) ----------------
@bp.route('/finance_summary')
def finance_summary():
    try:
        date_from = _str(request.args.get('date_from')) or None
        date_to = _str(request.args.get('date_to')) or None
        data, df, dt = metrics.build_finance_summary(date_from, date_to)
        total_income = round(sum(d['income_y'] for d in data), 2)
        total_cost = round(sum(d['cost_y'] for d in data), 2)
        return ok({'list': data, 'total_income_y': total_income, 'total_cost_y': total_cost,
                   'total_profit_y': round(total_income - total_cost, 2),
                   'date_from': df, 'date_to': dt, 'demo': metrics._LAST_SOURCE == 'demo',
                   'data_source': metrics._LAST_SOURCE,
                   'data_date': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')})
    except Exception as e:
        return err(e)

@bp.route('/finance_summary/export')
def finance_summary_export():
    try:
        date_from = _str(request.args.get('date_from')) or None
        date_to = _str(request.args.get('date_to')) or None
        data, df, dt = metrics.build_finance_summary(date_from, date_to)
        headers = ['代理商ID', '代理商名称', '收入合计(元)', '支出合计(元)', '毛利(元)', '产出投入比',
                   '销售金额(元)', '换电金额(元)', '电费支出(元)', '账单支出(元)', '账单收入(元)']
        rows = [[d['agency_id'], d['agency_name'], d['income_y'], d['cost_y'], d['profit_y'], d['ratio'],
                 d['sale_y'], d['order_y'], d['elec_y'], d['bill_out_y'], d['bill_in_y']] for d in data]
        return _export_excel(headers, rows, '财务收支汇总', '财务收支汇总_%s_%s.xlsx' % (df, dt))
    except Exception as e:
        return err(e)

# ---------------- 费用明细账单 ----------------
@bp.route('/expense_bills')
def expense_bills():
    try:
        agency_id = _int(request.args.get('agency_id'))
        keyword = _str(request.args.get('keyword'))
        date_from = _str(request.args.get('date_from')) or None
        date_to = _str(request.args.get('date_to')) or None
        page = max(1, _int(request.args.get('page'), 1))
        page_size = min(200, max(10, _int(request.args.get('page_size'), 50)))
        rows, total, df, dt = metrics.build_expense_bills(
            agency_id=agency_id or None, date_from=date_from, date_to=date_to,
            keyword=keyword or None, page=page, page_size=page_size,
            bill_id=_int(request.args.get('bill_id')) or None,
            business_id=_int(request.args.get('business_id')) or None,
            site_name=_str(request.args.get('site_name')) or None,
            merchant_name=_str(request.args.get('merchant_name')) or None,
            city=_str(request.args.get('city')) or None)
        return ok(rows, total=total, page=page, page_size=page_size,
                  date_from=df, date_to=dt, demo=metrics._LAST_SOURCE == 'demo',
                  data_source=metrics._LAST_SOURCE,
                  data_date=datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
    except Exception as e:
        return err(e)

@bp.route('/expense_bills/export')
def expense_bills_export():
    try:
        agency_id = _int(request.args.get('agency_id'))
        keyword = _str(request.args.get('keyword'))
        date_from = _str(request.args.get('date_from')) or None
        date_to = _str(request.args.get('date_to')) or None
        # 全量导出(一次性拉取, 避免逐页多次SQL)
        rows_all, total, df, dt = metrics.build_expense_bills(
            agency_id=agency_id or None, date_from=date_from, date_to=date_to,
            keyword=keyword or None, page=1, page_size=-1,
            bill_id=_int(request.args.get('bill_id')) or None,
            business_id=_int(request.args.get('business_id')) or None,
            site_name=_str(request.args.get('site_name')) or None,
            merchant_name=_str(request.args.get('merchant_name')) or None,
            city=_str(request.args.get('city')) or None)
        headers = ['账单ID', '账单日期', '结算日期', '业务类型', '业务ID', '费用名称', '费用类型', '费用值', '费用值类型',
                   '基准描述', '金额(税前/元)', '金额(税后/元)', '收入方类型', '收入方ID', '收入方名称',
                   '支出方类型', '支出方ID', '支出方名称', '代理商ID', '代理商名称']
        data = [[r['bill_id'], r['bill_date'], r['settle_date'], r['business_type'], r['business_id'], r['expense_name'],
                 r['expense_type'], r['expense_value'], r['expense_value_type'], r['standard_expense_desc'],
                 r['fee_y'], r['after_taxes_fee_y'], r['in_unit'], r['in_unit_id'], r['in_unit_name'],
                 r['out_unit'], r['out_unit_id'], r['out_unit_name'], r['agency_id'], r['agency_name']] for r in rows_all]
        return _export_excel(headers, data, '费用明细账单', '费用明细账单_%s_%s.xlsx' % (df, dt))
    except Exception as e:
        return err(e)

# ---------------- 总看板 ----------------
@bp.route('/overview')
def overview():
    try:
        date_from = _str(request.args.get('date_from')) or None
        date_to = _str(request.args.get('date_to')) or None
        data = metrics.build_overview(date_from, date_to)
        data['demo'] = metrics._LAST_SOURCE == 'demo'
        data['data_source'] = metrics._LAST_SOURCE
        data['data_date'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        return ok(data)
    except Exception as e:
        return err(e)

# ---------------- 财务看板(成本费用结构) ----------------
@bp.route('/finance_detail')
def finance_detail():
    try:
        date_from = _str(request.args.get('date_from')) or None
        date_to = _str(request.args.get('date_to')) or None
        data = metrics.build_finance_detail(date_from, date_to)
        data['demo'] = metrics._LAST_SOURCE == 'demo'
        data['data_source'] = metrics._LAST_SOURCE
        data['data_date'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        return ok(data)
    except Exception as e:
        return err(e)

@bp.route('/finance_detail/export')
def finance_detail_export():
    try:
        date_from = _str(request.args.get('date_from')) or None
        date_to = _str(request.args.get('date_to')) or None
        data = metrics.build_finance_detail(date_from, date_to)
        headers = ['分类', '项目', '金额(元)', '收支方向']
        rows = [[it['cate'], it['name'], it['value_y'], '收入' if it['kind'] == 'income' else '支出'] for it in data['items']]
        s = data['summary']
        rows.append(['合计', '总支出', s['total_expense_y'], '支出'])
        rows.append(['合计', '净利润(营收-总支出)', s['net_profit_y'], '利润'])
        return _export_excel(headers, rows, '财务看板', '财务看板成本费用_%s_%s.xlsx' % (data['date_from'], data['date_to']))
    except Exception as e:
        return err(e)
